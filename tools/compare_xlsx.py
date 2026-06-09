#!/usr/bin/env python3
"""A/B-сравнение XLSX-выгрузок по листу «Все контакты».

Два режима:
    python tools/compare_xlsx.py old.xlsx new.xlsx        # два файла
    python tools/compare_xlsx.py --dir OLD_DIR NEW_DIR    # две папки: сумма по всем XLSX
                                                           # + парные файлы по совпадению имён

Считает заполненность ключевых колонок и Δ. Колонки, которых нет в файле
(напр. «Имя Отчество» в старой версии), считаются 0.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

# Заголовок колонки в листе → человекочитаемая метрика. Считаем непустые ячейки.
_COLS = [
    ("ФИО (полностью)", "с ФИО"),
    ("Имя Отчество", "Имя Отчество"),
    ("Пол", "с полом (М/Ж)"),
    ("Должность (норм.)", "норм. должность"),
    ("Категория отрасли", "категория"),
    ("Личный email", "личный email"),
    ("Общий email", "общий email"),
    ("Личный телефон", "личный тел"),
    ("Общий телефон", "общий тел"),
    ("ИНН", "ИНН"),
    ("КПП", "КПП"),
    ("ОГРН", "ОГРН"),
    ("Соцсети", "соцсети"),
]

_SHEETS = ["Генеральные директора", "Финансовые директора",
           "Главные бухгалтеры", "Главные инженеры", "Остальные"]


def _nonempty(v) -> bool:
    if v is None:
        return False
    s = str(v).strip()
    # «Пол»: «?» считаем как НЕ определён
    return bool(s) and s != "?"


def metrics(path: str) -> dict:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["Все контакты"]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        wb.close()
        return {"всего строк": 0}
    header = [str(h) if h is not None else "" for h in rows[0]]
    data = rows[1:]
    m: dict[str, int] = {"всего строк": len(data)}
    for col, label in _COLS:
        if col in header:
            i = header.index(col)
            m[label] = sum(1 for r in data if i < len(r) and _nonempty(r[i]))
        else:
            m[label] = 0  # колонки нет в этой версии
    # счётчики листов
    for name in _SHEETS:
        if name in wb.sheetnames:
            cnt = wb[name].max_row - 1  # минус заголовок
            m[f"лист «{name}»"] = max(0, cnt)
    wb.close()
    return m


def _print_table(ma: dict, mb: dict) -> None:
    keys = list(dict.fromkeys(list(ma) + list(mb)))
    print(f"{'метрика':28} {'A':>9} {'B':>9} {'Δ':>9}")
    print("-" * 58)
    for k in keys:
        va, vb = ma.get(k, 0), mb.get(k, 0)
        d = vb - va
        mark = "  ←" if d else ""
        print(f"{k:28} {va:>9} {vb:>9} {d:>+9}{mark}")


def metrics_dir(d: str) -> tuple[dict, list[str]]:
    """Сумма метрик по всем *.xlsx в папке."""
    total: dict[str, int] = {}
    names: list[str] = []
    for f in sorted(Path(d).glob("*.xlsx")):
        try:
            m = metrics(str(f))
        except Exception as e:  # noqa: BLE001
            print(f"  ! пропуск {f.name}: {e}")
            continue
        names.append(f.name)
        for k, v in m.items():
            total[k] = total.get(k, 0) + v
    return total, names


def main(argv: list[str]) -> int:
    if argv and argv[0] == "--dir":
        if len(argv) < 3:
            print(__doc__)
            return 2
        old_dir, new_dir = argv[1], argv[2]
        ma, fa = metrics_dir(old_dir)
        mb, fb = metrics_dir(new_dir)
        print(f"\nA (старая): {old_dir}  ({len(fa)} xlsx)")
        print(f"B (новая):  {new_dir}  ({len(fb)} xlsx)\n")
        print("=== СУММА по всем файлам ===")
        _print_table(ma, mb)
        # Парные файлы по совпадению имён (если новые переименованы под старые id).
        common = sorted(set(fa) & set(fb))
        if common:
            print(f"\n=== Парные файлы ({len(common)}) — ключевые Δ ===")
            keys = ["всего строк", "с ФИО", "Имя Отчество", "с полом (М/Ж)",
                    "норм. должность", "личный email", "КПП"]
            print(f"{'файл':30} " + " ".join(f"{k[:9]:>9}" for k in keys))
            for name in common:
                a = metrics(str(Path(old_dir) / name))
                b = metrics(str(Path(new_dir) / name))
                deltas = " ".join(f"{(b.get(k,0)-a.get(k,0)):>+9}" for k in keys)
                print(f"{name[:30]:30} {deltas}")
        only_a = sorted(set(fa) - set(fb))
        only_b = sorted(set(fb) - set(fa))
        if only_a:
            print(f"\nтолько в A ({len(only_a)}): {', '.join(only_a)}")
        if only_b:
            print(f"только в B ({len(only_b)}): {', '.join(only_b)}")
        return 0

    if len(argv) < 2:
        print(__doc__)
        return 2
    a, b = argv[0], argv[1]
    print(f"\n{'A (старая):':<14}{a}")
    print(f"{'B (новая):':<14}{b}\n")
    _print_table(metrics(a), metrics(b))
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
